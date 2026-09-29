"""심층 분석 결과와 회차별 고정 2게임을 로컬에 저장한다."""

from datetime import datetime, timedelta
import hashlib
import json
from math import comb
import os
from pathlib import Path
import random
import tempfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

import lotto
import lotto_analysis as analysis


DRAW_SOURCE = "https://news.imbc.com/original/mbig/6659750_29041.html"
FACT_SOURCE = "https://factcheck.dhlottery.co.kr/"


def history_hash(draws):
    data = [[d.round, d.drawn_at.isoformat(), list(d.numbers), d.bonus] for d in draws]
    return hashlib.sha256(json.dumps(data, separators=(",", ":")).encode()).hexdigest()


def validate_record(record, target, fingerprint):
    try:
        tickets = record["tickets"]
        valid = (record["target_round"] == target and record["source_hash"] == fingerprint
                 and len(tickets) == 2 and tickets[0] != tickets[1])
        for ticket in tickets:
            valid = (valid and len(ticket) == 6 and len(set(ticket)) == 6
                     and all(type(n) is int and 1 <= n <= 45 for n in ticket)
                     and ticket == sorted(ticket))
        valid = valid and record["analysis"]["evaluation"]["selected"] in analysis.MODELS
        if not valid:
            raise ValueError("회차·데이터 지문·저장 번호 불일치")
    except (KeyError, TypeError, ValueError) as exc:
        raise lotto.LottoError(f"기존 주간 기록을 확인해야 합니다. 저장 번호를 덮어쓰지 않았습니다: {exc}") from exc
    return record


def freeze_record(path, record):
    """완성된 파일만 공개하고, 동시 실행해도 먼저 저장한 회차 기록을 보존한다."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent,
                                     suffix=".tmp", delete=False) as handle:
        temporary = Path(handle.name)
        json.dump(record, handle, ensure_ascii=False, indent=2, allow_nan=False)
    try:
        try:
            os.link(temporary, path)
        except FileExistsError:
            pass
        saved = json.loads(path.read_text(encoding="utf-8"))
        return validate_record(saved, record["target_round"], record["source_hash"])
    finally:
        temporary.unlink(missing_ok=True)


def weekly(path=lotto.DEFAULT_FILE, offline=False, log=print):
    path = Path(path).resolve()
    if not offline:
        log("공식 API의 최신 확정 회차를 확인합니다.")
        official = lotto.fetch_draw()
        lotto.update_history(path, official.round, log)
    draws = lotto.load_draws(path)
    if any(d.bonus is None for d in draws):
        raise lotto.LottoError("보너스 번호가 빠진 회차가 있습니다. 업데이트를 먼저 실행하세요.")
    if not offline:
        if draws[-1] != official:
            raise lotto.LottoError("엑셀 마지막 회차와 공식 최신 결과가 다릅니다. 기존 값을 확인하세요.")
        if datetime.now(lotto.KST).date() > official.drawn_at + timedelta(days=7):
            raise lotto.LottoError("공식 응답이 이전 주에 머물러 있습니다. 최신 결과를 다시 확인한 뒤 실행하세요.")
    target = draws[-1].round + 1
    fingerprint = history_hash(draws)
    record_path = path.parent / "weekly" / f"{target}.json"
    if record_path.exists():
        try:
            record = json.loads(record_path.read_text(encoding="utf-8"))
        except (ValueError, OSError) as exc:
            raise lotto.LottoError(f"저장된 주간 기록을 읽지 못했습니다. 재추출하지 않습니다: {exc}") from exc
        validate_record(record, target, fingerprint)
        log(f"{target}회에 처음 저장한 2게임을 그대로 불러옵니다.")
    else:
        result = analysis.analyze(draws, log)
        selected = result["evaluation"]["selected"]
        tickets = analysis.sample_two(result["evaluation"]["next_probabilities"][selected], random.SystemRandom())
        record = {"version": analysis.VERSION, "target_round": target, "source_round": draws[-1].round,
                  "source_date": draws[-1].drawn_at.isoformat(), "source_hash": fingerprint,
                  "planned_date": (draws[-1].drawn_at + timedelta(days=7)).isoformat(),
                  "created_at": datetime.now(lotto.KST).isoformat(timespec="seconds"),
                  "online_verified_at_creation": not offline, "tickets": tickets, "analysis": result}
        if history_hash(lotto.load_draws(path)) != fingerprint:
            raise lotto.LottoError("분석 중 이력이 바뀌었습니다. 다시 실행하세요. 주간 번호는 저장하지 않았습니다.")
        record = freeze_record(record_path, record)
    report, summary = export_reports(record, path.parent)
    log(f"{target}회 2게임 · 추첨 예정 {record['planned_date']}")
    for i, ticket in enumerate(record["tickets"], 1):
        log(f"  게임 {i}: {'  '.join(f'{n:02}' for n in ticket)}")
    log(record["analysis"]["evaluation"]["reason"])
    log(f"주간 결과: {report}")
    log(f"심층 분석 설명: {summary}")
    return {"record": record, "record_path": record_path, "report": report, "summary": summary}


def export_reports(record, directory):
    directory = Path(directory) / "reports"
    directory.mkdir(parents=True, exist_ok=True)
    report = directory / f"weekly-{record['target_round']}.xlsx"
    summary = directory / f"weekly-{record['target_round']}.md"
    evaluation = record["analysis"]["evaluation"]
    diagnosis = record["analysis"]["diagnostics"]
    if not report.exists():
        book = Workbook()
        notes = book.active
        notes.title = "주간추천"
        note_rows = [("항목", "내용"), ("대상회차", record["target_round"]),
                     ("추첨 예정일", record["planned_date"]), ("기준 데이터", f"1~{record['source_round']}회"),
                     ("선택 방식", analysis.MODELS[evaluation["selected"]]), ("선택 이유", evaluation["reason"]),
                     ("게임 1", " · ".join(f"{n:02}" for n in record["tickets"][0])),
                     ("게임 2", " · ".join(f"{n:02}" for n in record["tickets"][1])),
                     ("저장 정책", "해당 회차에 처음 저장한 서로 다른 2게임을 유지합니다. 재실행해도 바뀌지 않습니다."),
                     ("추첨 방식", "당첨번호는 추첨기와 추첨볼로 결정됩니다. 컴퓨터 난수 시드 복원 대상이 아닙니다."),
                     ("데이터 해석", "엑셀의 번호1~6은 오름차순 번호입니다. 실제 공 배출 순서가 아닙니다."),
                     ("두 게임의 1등 확률", f"공정한 추첨: 2 / {comb(45, 6):,} = 1 / {comb(45, 6) // 2:,}"),
                     ("검정 해석", "보정 p<0.05는 무작위 기준과의 차이를 탐색하는 기준입니다. 원인이나 미래 예측력을 증명하지 않습니다. 비유의도 공정성의 증명은 아닙니다."),
                     ("확률오차", "Brier 점수는 번호 45개의 출현 여부와 예측확률의 제곱오차 평균입니다. 낮을수록 좋습니다. 6개 조합의 당첨 확률을 직접 검증한 값은 아닙니다."),
                     ("표본 분리", f"선택: {evaluation['validation_start']}~{evaluation['validation_end']}회 / 분리 평가: {evaluation['test_start']}~{evaluation['test_end']}회"),
                     ("평가 방법", "후보는 선택 구간 점수로만 정합니다. 분리 평가에서 탈락하면 다른 후보를 고르지 않고 균등 확률을 사용합니다."),
                     ("오차 개선 구간", "13회 단위 원형 블록 부트스트랩 4,000회, 근사 95% 구간. 양수이면 균등 예측보다 오차가 작다는 뜻입니다."),
                     ("확률 모델", "누적·최근 빈도에는 균등 확률 52회분을 사전정보로 더합니다. 최근 가중 평균은 반감기 26회입니다. 고정 설정을 평가 결과에 맞춰 바꾸지 않습니다."),
                     ("2게임 과거실험", "회차별 고정 시험용 난수로 2게임씩 추출한 한 번의 참고 실험입니다. 이 결과로 모델을 선택하지 않습니다."),
                     ("가중 추출", "통계 방식 채택 시 예측값을 추출 가중치로 사용합니다. 이 가중치가 실제 포함확률이나 당첨확률과 같지는 않습니다."),
                     ("자료 출처", lotto.API), ("추첨 현장", DRAW_SOURCE), ("공식 설명", FACT_SOURCE),
                     ("최초 생성 시각", record["created_at"]), ("온라인 확인 후 생성", record["online_verified_at_creation"]),
                     ("이력 지문", record["source_hash"]), ("분석 버전", record["version"])]
        for row in note_rows:
            notes.append(row)
        tickets = book.create_sheet("추천번호")
        tickets.append(["회차", "추첨 예정일", "게임", *lotto.HEADERS[2:8]])
        for i, ticket in enumerate(record["tickets"], 1):
            tickets.append([record["target_round"], record["planned_date"], i, *ticket])
        tests = book.create_sheet("무작위성진단")
        tests.append(["분류", "대상", "관측값", "기준값", "원래 p", "전체 Holm 보정 p", "보정 후 p<0.05", "검정 방법"])
        for row in sorted(diagnosis["tests"], key=lambda item: item["p"]):
            tests.append([row["group"], row["name"], row["observed"], row["expected"], row["p"], row["adjusted_p"], row["significant"], row["method"]])
        distributions = book.create_sheet("조합분포")
        distributions.append(["분류", "구간", "실제 횟수", "무작위 기준 기대 횟수"])
        for row in diagnosis["distributions"]:
            distributions.append([row["group"], row["bin"], row["observed"], row["expected"]])
        models = book.create_sheet("모델비교")
        models.append(["모델", "선택 구간 Brier", "분리 평가 Brier", "분리 평가 로그손실", "선택 구간 보정 p",
                       "분리 평가 개선", "개선 95% 하한", "개선 95% 상한", "참고 2게임 평균 일치", "선택 구간 최우수 후보", "최종 사용"])
        examples = book.create_sheet("2게임과거실험")
        examples.append(["모델", "시작회차", "종료회차", "전체 게임 수", "평균 일치", "1등", "2등", "3등", "4등", "5등"])
        for result in evaluation["models"]:
            models.append([result["label"], result["validation_brier"], result["test_brier"], result["test_log_loss"],
                           result["validation"]["adjusted_p"], result["test"]["gain"], result["test"]["ci_low"], result["test"]["ci_high"],
                           result["illustrative_mean_hits"], result["model"] == evaluation["candidate"], result["model"] == evaluation["selected"]])
            examples.append([result["label"], evaluation["test_start"], evaluation["test_end"], analysis.HOLDOUT * 2,
                             result["illustrative_mean_hits"], *[result["illustrative_ranks"][str(n)] for n in range(1, 6)]])
        probabilities = book.create_sheet("다음회차분석값")
        probabilities.append(["번호", *analysis.MODELS.values()])
        for number in range(45):
            probabilities.append([number + 1, *[evaluation["next_probabilities"][model][number] for model in analysis.MODELS]])
        frequency = book.create_sheet("번호빈도")
        frequency.append(["번호", "관측 출현 횟수", "기대 출현 횟수", "보정 p"])
        for row in diagnosis["tests"][:45]:
            frequency.append([int(row["name"]), row["observed"], row["expected"], row["adjusted_p"]])
        chart = BarChart()
        chart.title = "전체 번호 빈도와 무작위 기준"
        chart.add_data(Reference(frequency, min_col=2, max_col=3, min_row=1, max_row=46), titles_from_data=True)
        chart.set_categories(Reference(frequency, min_col=1, min_row=2, max_row=46))
        chart.width, chart.height = 30, 12
        frequency.add_chart(chart, "F2")
        lotto.style_report(book)
        notes.column_dimensions["B"].width = 110
        for index in range(2, notes.max_row + 1):
            notes.row_dimensions[index].height = 44
        try:
            lotto.save_atomic(book, report)
        finally:
            book.close()
    if not summary.exists():
        summary.write_text(markdown_report(record), encoding="utf-8")
    return report, summary


def markdown_report(record):
    result = record["analysis"]
    diagnosis, evaluation = result["diagnostics"], result["evaluation"]
    numbers = diagnosis["tests"][:45]
    low, high = min(numbers, key=lambda row: row["observed"]), max(numbers, key=lambda row: row["observed"])
    lines = [f"# {record['target_round']}회 주간 2게임 및 심층 분석", "",
             f"기준: 1~{record['source_round']}회, 마지막 추첨 {record['source_date']}. 다음 추첨 예정일: {record['planned_date']}.", "",
             *[f"- 게임 {i}: **{' · '.join(f'{n:02}' for n in ticket)}**" for i, ticket in enumerate(record["tickets"], 1)], "",
             f"선택 방식: **{analysis.MODELS[evaluation['selected']]}**. {evaluation['reason']}", "",
             "같은 회차의 번호는 최초 기록을 유지합니다. 컴퓨터 난수의 시드를 알아낸 결과나 당첨 보장이 아닙니다.", "",
             "## 추첨 원리와 데이터 한계", "",
             f"로또 당첨번호는 추첨기와 추첨볼로 결정됩니다. 구매 시 자동선택되는 번호와 당첨 추첨은 서로 다른 과정입니다. [MBC 추첨 현장]({DRAW_SOURCE}), [동행복권 설명]({FACT_SOURCE}).", "",
             "엑셀의 여섯 번호는 오름차순이며 실제 공 배출 순서가 없습니다. 열별 큰 수·작은 수의 차이나 정렬된 숫자열의 상관을 난수 생성기의 흔적으로 해석할 수 없습니다. 원시 난수 비트열도 없으므로 시드 복원이나 비트열 난수성 인증을 하지 않습니다.", "",
             "## 실제 데이터 진단", "",
             f"- 회차 {diagnosis['draws']:,}개, 동일한 6개 조합 반복 {diagnosis['duplicate_combinations']}건, 7일 간격이 아닌 추첨일 {diagnosis['nonweekly_dates']}건.",
             f"- 한 번호의 기대 출현은 {diagnosis['number_expected']:.2f}회. 최다 {high['name']}번 {high['observed']:.0f}회, 최소 {low['name']}번 {low['observed']:.0f}회.",
             f"- 특정 번호쌍의 동시 출현 확률은 1/66, 기대 출현은 {diagnosis['pair_expected']:.2f}회.",
             f"- 번호 합계 평균 {diagnosis['sum_mean']:.3f}, 이론 평균 138. 직전 회차와 중복 평균 {diagnosis['overlap_mean']:.4f}개, 이론 평균 0.8개.",
             f"- 번호 45개, 번호쌍 990개, 전후반 변화 45개, 회차 간격 10개, 조합 분포 3개의 **총 {len(diagnosis['tests']):,}개 검정**을 하나의 집합으로 Holm 보정했습니다.",
             f"- 보정 p<0.05인 항목: **{diagnosis['significant_count']}개**. 아래는 원래 p가 작은 순서의 진단이며, 보정값으로 판단해야 합니다.", "",
             "|분류|대상|관측값|기준값|원래 p|전체 보정 p|", "|---|---|---:|---:|---:|---:|"]
    for row in sorted(diagnosis["tests"], key=lambda row: row["p"])[:12]:
        lines.append(f"|{row['group']}|{row['name']}|{row['observed']:.4f}|{row['expected']:.4f}|{row['p']:.6f}|{row['adjusted_p']:.6f}|")
    lines += ["", "번호와 번호쌍은 양측 정확 이항검정, 전후반은 Fisher 정확검정입니다. 회차 간 평균 중복은 비복원 추출 분산을 이용한 큰 표본 정규 근사입니다. 홀짝·합계·연속수는 정확한 조합 확률을 기준으로 카이제곱 근사 검정을 합니다. ‘조합 분포’ 행의 관측값은 검정통계량이고 기준값은 자유도입니다.", "",
              "유의한 항목이 있어도 원인이나 예측력을 뜻하지 않으며, 유의하지 않다는 결과도 공정성의 증명이 아닙니다. 보너스는 빈도 진단에서 제외하고 2등 판정에만 사용합니다.", "",
              "## 과거 결과를 이용한 예측 성능 평가", "",
              f"초기 {analysis.WARMUP}회 이후 매 회차 그 이전 데이터만 사용합니다. 후보 선택 구간은 {evaluation['validation_start']}~{evaluation['validation_end']}회, 선택에서 제외한 분리 평가 구간은 {evaluation['test_start']}~{evaluation['test_end']}회입니다. 분리 평가 중에도 이미 끝난 회차는 다음 회차의 학습에 사용할 수 있습니다.", "",
              "번호 45개 각각의 출현확률을 예측하고 Brier 점수(평균 제곱오차)를 계산합니다. 낮을수록 좋으며 균등 기준은 각 번호 6/45입니다. 이는 1등 조합 자체의 예측 성능과 같지 않습니다.", "",
              "|방식|선택 구간 Brier|분리 평가 Brier|분리 평가 개선|개선의 근사 95% 구간|", "|---|---:|---:|---:|---|" ]
    for row in evaluation["models"]:
        lines.append(f"|{row['label']}|{row['validation_brier']:.8f}|{row['test_brier']:.8f}|{row['test']['gain']:+.8f}|[{row['test']['ci_low']:+.8f}, {row['test']['ci_high']:+.8f}]|")
    lines += ["", f"선택 구간에서 정한 후보는 **{analysis.MODELS[evaluation['candidate']]}**입니다. 후보 6개의 선택 구간 p를 Holm 보정하고, 그 후보 하나만 분리 평가에서 재확인합니다. 두 구간 모두 개선의 95% 하한이 양수이고 관련 p<0.05일 때만 통계 가중치를 채택합니다. 나머지 후보의 분리 평가 점수는 설명용이며 다른 후보를 재선택하지 않습니다.", "",
              f"개선 구간과 p는 {evaluation['bootstrap_block']}회 단위 원형 블록 부트스트랩 {evaluation['bootstrap_samples']:,}회를 이용한 근사치입니다. 블록 길이·과거 표본·시장/장비 변화 등에 영향을 받을 수 있으며 미래 성능 보장이 아닙니다.", "",
              "누적·최근 26/52/104회 빈도는 균등 확률 52회분으로 완화했습니다. 지수 가중 평균의 반감기는 26회입니다. 전이 방식은 직전 여섯 번호 각각이 나온 뒤 다음 회차의 번호 출현 비율을 같은 방식으로 완화해 평균합니다. 날짜·회차를 난수 시드로 맞추거나 과거 성적에 맞춰 설정을 탐색하지 않았습니다.", "",
              "엑셀의 ‘2게임과거실험’은 각 평가 회차에 고정 시험용 난수로 2게임씩 추출한 참고 실험입니다. 번호 겹침과 당첨 횟수의 우연한 변동이 크므로 이 성적으로 모델을 선택하지 않습니다. 통계 방식의 예측값을 추출 가중치로 쓰더라도 실제 번호 포함확률과 정확히 같지는 않습니다.", "",
              "## 이번 회차 생성과 보관", "",
              "운영체제 난수로 선택한 방식의 가중치에 따라 여섯 번호를 중복 없이 뽑고, 서로 다른 두 조합을 저장합니다. 균등 방식을 선택하면 모든 조합이 같은 확률로 생성됩니다. 홀짝·합계·연속수·최근 출현 여부로 조합을 강제로 제외하지 않습니다.", "",
              f"공정한 추첨에서 두 개의 서로 다른 게임 중 1등이 있을 확률은 2/{comb(45,6):,} = 1/{comb(45,6)//2:,}입니다. 분석으로 이 확률이 커졌다는 근거는 없습니다.", "",
              f"기록: `weekly/{record['target_round']}.json`, 생성 시각: {record['created_at']}, 분석 버전: {record['version']}. 최초 생성 시 온라인 확인: {record['online_verified_at_creation']}.", "",
              f"데이터 지문: `{record['source_hash']}`. 실제 번호·추첨일·보너스가 바뀌면 기존 추천을 자동 교체하지 않고 확인 오류를 표시합니다.", "",
              "검정 구현 참고: [SciPy 이항검정](https://docs.scipy.org/doc/scipy-1.17.0/reference/generated/scipy.stats.binomtest.html), [Fisher 정확검정](https://docs.scipy.org/doc/scipy-1.17.0/reference/generated/scipy.stats.fisher_exact.html), [카이제곱검정](https://docs.scipy.org/doc/scipy-1.17.0/reference/generated/scipy.stats.chisquare.html).", ""]
    return "\n".join(lines)
