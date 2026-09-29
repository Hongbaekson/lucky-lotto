"""주간 추천의 시간순 검증, 기록 보존, 결과 파일을 검사한다."""

from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import date, timedelta
import json
from pathlib import Path
import random
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import Workbook, load_workbook

import lotto
import lotto_analysis as analysis
import lotto_weekly as weekly


def history(count=624):
    rng = random.Random(37)
    draws = []
    for number in range(1, count + 1):
        numbers = tuple(sorted(rng.sample(range(1, 46), 6)))
        bonus = next(n for n in range(1, 46) if n not in numbers)
        draws.append(lotto.Draw(number, date(2002, 12, 7) + timedelta(weeks=number - 1),
                                numbers, bonus, 1000000, 10, 100000000))
    return draws


class WeeklyChecks(unittest.TestCase):
    def test_diagnostics_use_correct_sampling_probabilities_and_multiple_testing(self):
        result = analysis.diagnostics(history())
        self.assertEqual(len(result["tests"]), 1093)
        self.assertAlmostEqual(sum(row["expected"] for row in result["tests"][:45]), 624 * 6)
        pairs = [row for row in result["tests"] if row["group"] == "번호쌍 빈도"]
        self.assertEqual(sum(row["observed"] for row in pairs), 624 * 15)
        self.assertAlmostEqual(pairs[0]["expected"], 624 / 66)
        for group in {row["group"] for row in result["distributions"]}:
            rows = [row for row in result["distributions"] if row["group"] == group]
            self.assertEqual(sum(row["observed"] for row in rows), 624)
            self.assertAlmostEqual(sum(row["expected"] for row in rows), 624)
        np.testing.assert_allclose(analysis.holm([0.01, 0.04, 0.03]), [0.03, 0.06, 0.06])
        self.assertTrue(all(row["p"] <= row["adjusted_p"] <= 1 for row in result["tests"]))

    def test_forecasts_cannot_see_the_target_or_future_rows(self):
        matrix = analysis.indicators(history(320))
        predictions = dict(analysis.forecasts(matrix))
        for cutoff in (208, 235, 300):
            prefix = dict(analysis.forecasts(matrix[:cutoff]))[cutoff]
            modified = matrix.copy()
            modified[cutoff:] = np.roll(modified[cutoff:], 7, axis=1)
            changed = dict(analysis.forecasts(modified))[cutoff]
            for model in analysis.MODELS:
                np.testing.assert_array_equal(predictions[cutoff][model], prefix[model])
                np.testing.assert_array_equal(predictions[cutoff][model], changed[model])
                self.assertAlmostEqual(float(prefix[model].sum()), 6)
                self.assertTrue(np.all((prefix[model] > 0) & (prefix[model] < 1)))

    def test_holdout_outcomes_cannot_change_the_candidate(self):
        draws = history()
        changed = [d if d.round <= 416 else lotto.Draw(d.round, d.drawn_at, (1, 2, 3, 4, 5, 6), 7, d.prize, d.winners, d.sales)
                   for d in draws]
        with patch.object(analysis, "BOOTSTRAPS", 255):
            before = analysis.evaluate(draws, log=lambda _: None)
            after = analysis.evaluate(changed, log=lambda _: None)
        self.assertEqual(before["candidate"], after["candidate"])
        self.assertEqual((before["validation_start"], before["validation_end"], before["test_start"], before["test_end"]),
                         (209, 416, 417, 624))
        for a, b in zip(before["models"], after["models"]):
            self.assertEqual(a["validation_brier"], b["validation_brier"])
            self.assertEqual(a["validation"], b["validation"])
        self.assertAlmostEqual(before["models"][0]["test_brier"], (6 / 45) * (39 / 45))

    def test_first_saved_two_games_win_even_with_concurrent_writers(self):
        record = {"target_round": 625, "source_hash": "example",
                  "tickets": [[1, 2, 3, 4, 5, 6], [7, 8, 9, 10, 11, 12]],
                  "analysis": {"evaluation": {"selected": "uniform"}}}
        other = deepcopy(record)
        other["tickets"] = [[20, 21, 22, 23, 24, 25], [30, 31, 32, 33, 34, 35]]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "weekly" / "625.json"
            with ThreadPoolExecutor(max_workers=2) as pool:
                saved = list(pool.map(lambda value: weekly.freeze_record(path, value), [record, other]))
            self.assertEqual(saved[0], saved[1])
            self.assertIn(saved[0], [record, other])
            self.assertEqual(json.loads(path.read_text(encoding="utf-8")), saved[0])
            self.assertEqual(list(path.parent.glob("*.tmp")), [])
            with self.assertRaises(lotto.LottoError):
                weekly.validate_record(saved[0], 625, "different history")

    def test_weekly_workflow_saves_two_and_reuses_the_record_without_reanalysis(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "lotto.xlsx"
            book = Workbook()
            book.active.append(lotto.HEADERS)
            for draw in history():
                book.active.append(draw.cells())
            book.save(path)
            book.close()
            original = path.read_bytes()
            with patch.object(analysis, "BOOTSTRAPS", 255):
                first = weekly.weekly(path, offline=True, log=lambda _: None)
            self.assertEqual(path.read_bytes(), original)
            saved_bytes = first["record_path"].read_bytes()
            with patch.object(analysis, "analyze", side_effect=AssertionError("unexpected reanalysis")):
                again = weekly.weekly(path, offline=True, log=lambda _: None)
            self.assertEqual(first["record"], again["record"])
            self.assertEqual(first["record_path"].read_bytes(), saved_bytes)
            tickets = first["record"]["tickets"]
            self.assertEqual(len(tickets), 2)
            self.assertNotEqual(tickets[0], tickets[1])
            for ticket in tickets:
                self.assertEqual(len(set(ticket)), 6)
                self.assertEqual(ticket, sorted(ticket))
                self.assertTrue(all(1 <= n <= 45 for n in ticket))
            report = load_workbook(first["report"])
            self.assertEqual(report["추천번호"].max_row, 3)
            self.assertEqual(report["무작위성진단"].max_row, 1094)
            self.assertEqual(report["모델비교"].max_row, 8)
            report.close()
            changed = load_workbook(path)
            changed.active["J2"] = 2000000  # 분석과 무관한 금액 변경은 주간 번호를 바꾸지 않는다.
            changed.save(path)
            changed.close()
            with patch.object(analysis, "analyze", side_effect=AssertionError("unexpected reanalysis")):
                self.assertEqual(weekly.weekly(path, offline=True, log=lambda _: None)["record"], first["record"])


if __name__ == "__main__":
    unittest.main()
