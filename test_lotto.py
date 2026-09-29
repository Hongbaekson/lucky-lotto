"""실제 파일을 건드리지 않는 데이터 보존/추천 검증: python -m unittest -v"""

from copy import deepcopy
from datetime import date, timedelta
from pathlib import Path
import random
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

import lotto


class LottoChecks(unittest.TestCase):
    def draw(self, number):
        return lotto.Draw(number, date(2002, 12, 7) + timedelta(weeks=number - 1),
                          (1, 8, 15, 22, 29, 36), 43, 1234567890, 7, 90000000000)

    def make_file(self, directory):
        path = Path(directory) / "lotto.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "로또당첨번호"
        sheet.append(lotto.HEADERS)
        sheet.append(self.draw(1).cells())
        sheet.append(self.draw(2).cells()[:8])  # 기존 마지막 회차에 보너스/금액 누락
        sheet["B2"].number_format = sheet["B3"].number_format = "yyyy-mm-dd"
        sheet["A3"].fill = PatternFill("solid", fgColor="FFD700")
        sheet["A100"].fill = PatternFill("solid", fgColor="EEEEEE")  # 값 없는 서식 행
        sheet.column_dimensions["B"].width = 18
        sheet["N2"] = "보존할 메모"
        book.save(path)
        book.close()
        return path

    def test_update_preserves_existing_values_and_styles_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.make_file(directory)
            original = path.read_bytes()
            with patch("lotto.fetch_draw", side_effect=self.draw), patch("lotto.time.sleep"):
                result = lotto.update_history(path, 4, log=lambda _: None)
            self.assertEqual((result["added"], result["repaired"]), (2, 1))
            self.assertEqual(result["backup"].read_bytes(), original)
            self.assertEqual(lotto.load_draws(path), [self.draw(n) for n in range(1, 5)])
            book = load_workbook(path)
            sheet = book.active
            self.assertEqual(sheet["N2"].value, "보존할 메모")
            self.assertEqual(sheet["A4"]._style, sheet["A3"]._style)
            self.assertEqual(sheet["B5"].number_format, "yyyy-mm-dd")
            self.assertEqual(sheet.column_dimensions["B"].width, 18)
            self.assertIsNone(sheet["A100"].value)
            book.close()
            saved = path.read_bytes()
            with patch("lotto.fetch_draw") as fetch:
                again = lotto.update_history(path, 4, log=lambda _: None)
                fetch.assert_not_called()
            self.assertEqual(again["added"], 0)
            self.assertEqual(path.read_bytes(), saved)

    def test_conflict_or_api_failure_never_saves_partial_results(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.make_file(directory)
            original = path.read_bytes()
            conflict = lotto.Draw(2, self.draw(2).drawn_at, (2, 8, 15, 22, 29, 36), 43, 1, 1, 1)
            with patch("lotto.fetch_draw", return_value=conflict):
                with self.assertRaises(lotto.LottoError):
                    lotto.update_history(path, 4, log=lambda _: None)
            self.assertEqual(path.read_bytes(), original)
            with patch("lotto.fetch_draw", side_effect=[self.draw(2), self.draw(3), lotto.LottoError("API 실패")]), patch("lotto.time.sleep"):
                with self.assertRaises(lotto.LottoError):
                    lotto.update_history(path, 4, log=lambda _: None)
            self.assertEqual(path.read_bytes(), original)
            self.assertFalse((Path(directory) / "backups").exists())

    def test_official_api_mapping_and_wrong_round_rejection(self):
        payload = {"data": {"list": [{"ltEpsd": 1243, "ltRflYmd": "20260926",
                    "tm1WnNo": 9, "tm2WnNo": 18, "tm3WnNo": 24, "tm4WnNo": 38,
                    "tm5WnNo": 43, "tm6WnNo": 44, "bnsWnNo": 35,
                    "rnk1WnAmt": 2592525282, "rnk1WnNope": 12,
                    "rlvtEpsdSumNtslAmt": 64463211121, "wholEpsdSumNtslAmt": 128926419000}]}}
        draw = lotto.parse_response(payload, 1243)
        self.assertEqual(draw.numbers, (9, 18, 24, 38, 43, 44))
        self.assertEqual(draw.sales, 128926419000)
        self.assertEqual(draw.prize, 2592525282)
        with self.assertRaises(lotto.LottoError):
            lotto.parse_response(payload, 1244)
        invalid = deepcopy(payload)
        invalid["data"]["list"][0]["bnsWnNo"] = 9
        with self.assertRaises(lotto.LottoError):
            lotto.parse_response(invalid, 1243)

    def test_recommendations_are_valid_reproducible_and_bonus_is_excluded(self):
        history = [self.draw(n) for n in range(1, 60)]
        for mode in lotto.MODES:
            tickets = lotto.recommend(history, 100, mode, rng=random.Random(7))
            self.assertEqual(tickets, lotto.recommend(history, 100, mode, rng=random.Random(7)))
            self.assertEqual(len(set(tickets)), 100)
            for ticket in tickets:
                self.assertEqual(len(set(ticket)), 6)
                self.assertEqual(ticket, tuple(sorted(ticket)))
                self.assertTrue(all(1 <= n <= 45 for n in ticket))
        stats = lotto.statistics(history)
        self.assertEqual(stats[42]["total"], 0)
        self.assertEqual(stats[0]["total"], 59)
        self.assertEqual(lotto.rank((1, 8, 15, 22, 29, 43), history[-1]), (5, 2))

    def test_backtest_never_reads_target_or_future_draws(self):
        history = [self.draw(n) for n in range(1, 61)]
        seen = []
        original_recommend = lotto.recommend

        def tracked(training, *args, **kwargs):
            seen.append(training[-1].round)
            return original_recommend(training, *args, **kwargs)

        with patch("lotto.recommend", side_effect=tracked):
            results = lotto.backtest(history, rounds=5, games=3, window=10)
        self.assertEqual(seen, list(range(55, 60)) * 3)
        self.assertEqual(results, lotto.backtest(history, rounds=5, games=3, window=10))
        for result in results:
            self.assertEqual((result["start"], result["end"], result["tickets"]), (56, 60, 15))
            self.assertEqual(sum(result["matches"].values()), 15)
            self.assertEqual(sum(result["ranks"].values()), 15)


if __name__ == "__main__":
    unittest.main()
