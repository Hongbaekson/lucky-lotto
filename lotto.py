"""동행복권 데이터 갱신, 설명 가능한 번호 추천, 시간순 과거 검증."""

from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time as dt_time, timedelta, timezone
import hashlib
import json
import os
from pathlib import Path
import random
import shutil
import tempfile
import time
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path(__file__).resolve().parent
DEFAULT_FILE = BASE / "lotto.xlsx"
API = "https://www.dhlottery.co.kr/lt645/selectPstLt645Info.do"
KST = timezone(timedelta(hours=9))
HEADERS = ("회차", "추첨일", "번호1", "번호2", "번호3", "번호4", "번호5", "번호6",
           "보너스", "1등당첨금", "1등당첨자수", "총판매금액")
MODES = {"random": "균등 무작위", "frequency": "전체 빈도 참고", "recent": "최근 빈도 참고"}
NOTICE = "과거 출현 빈도는 미래 당첨 확률을 높인다는 근거가 없습니다. 모든 조합의 1등 확률은 1/8,145,060입니다."


class LottoError(ValueError):
    """사용자에게 그대로 표시할 수 있는 데이터/입력 오류."""


def integer(value, label, minimum=0, optional=False):
    if value is None and optional:
        return None
    if isinstance(value, bool):
        raise LottoError(f"{label}: 정수가 필요합니다.")
    try:
        result = int(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise LottoError(f"{label}: 정수가 필요합니다 ({value!r}).") from exc
    if str(value).strip() != str(result) and value != result:
        raise LottoError(f"{label}: 정수가 필요합니다 ({value!r}).")
    if result < minimum:
        raise LottoError(f"{label}: {minimum} 이상이어야 합니다.")
    return result


@dataclass(frozen=True)
class Draw:
    round: int
    drawn_at: date
    numbers: tuple[int, ...]
    bonus: int | None
    prize: int | None
    winners: int | None
    sales: int | None

    def cells(self):
        return (self.round, datetime.combine(self.drawn_at, dt_time()), *self.numbers,
                self.bonus, self.prize, self.winners, self.sales)


def parse_draw(values, complete=False):
    if len(values) != 12:
        raise LottoError("당첨 데이터는 12개 열이어야 합니다.")
    number = integer(values[0], "회차", 1)
    drawn_at = values[1]
    if isinstance(drawn_at, datetime):
        drawn_at = drawn_at.date()
    if not isinstance(drawn_at, date) or drawn_at > datetime.now(KST).date():
        raise LottoError(f"{number}회: 유효한 과거 추첨일이 필요합니다.")
    numbers = tuple(integer(n, f"{number}회 당첨번호", 1) for n in values[2:8])
    if len(set(numbers)) != 6 or max(numbers) > 45 or numbers != tuple(sorted(numbers)):
        raise LottoError(f"{number}회: 당첨번호는 1~45의 서로 다른 오름차순 숫자 6개여야 합니다.")
    bonus = integer(values[8], f"{number}회 보너스", 1, optional=not complete)
    if bonus is not None and (bonus > 45 or bonus in numbers):
        raise LottoError(f"{number}회: 보너스 번호가 잘못되었습니다.")
    amounts = [integer(v, f"{number}회 {HEADERS[i]}", optional=not complete)
               for i, v in enumerate(values[9:], 9)]
    return Draw(number, drawn_at, numbers, bonus, *amounts)


def parse_response(payload, requested=None):
    try:
        items = payload["data"]["list"]
        if not isinstance(items, list) or len(items) != 1:
            raise LottoError("공식 API에 해당 회차의 확정 결과가 없습니다.")
        item = items[0]
        draw = parse_draw((item["ltEpsd"], datetime.strptime(item["ltRflYmd"], "%Y%m%d"),
                           *(item[f"tm{i}WnNo"] for i in range(1, 7)), item["bnsWnNo"],
                           item["rnk1WnAmt"], item["rnk1WnNope"], item["wholEpsdSumNtslAmt"]),
                          complete=True)
    except (KeyError, TypeError, ValueError) as exc:
        raise LottoError(f"공식 API 응답을 확인할 수 없습니다: {exc}") from exc
    if requested is not None and draw.round != requested:
        raise LottoError(f"요청한 {requested}회 대신 {draw.round}회가 반환되어 저장을 중단했습니다.")
    return draw


def fetch_draw(number=None):
    if number is not None:
        number = integer(number, "조회 회차", 1)
    url = API if number is None else f"{API}?srchLtEpsd={number}"
    request = Request(url, headers={"Accept": "application/json", "User-Agent": "LuckyLotto/1.0",
                                   "Referer": "https://www.dhlottery.co.kr/lt645/result"})
    for attempt in range(3):
        try:
            with urlopen(request, timeout=20) as response:
                raw = response.read(1_000_000)
            try:
                payload = json.loads(raw)
            except (ValueError, UnicodeError) as exc:
                raise LottoError("API에서 JSON이 아닌 응답을 받았습니다. 접속 제한 또는 점검 여부를 확인하세요.") from exc
            return parse_response(payload, number)
        except (URLError, TimeoutError, OSError) as exc:
            retryable = not isinstance(exc, HTTPError) or exc.code in (429, 500, 502, 503, 504)
            if attempt == 2 or not retryable:
                raise LottoError(f"공식 API 연결 실패: {exc}") from exc
            time.sleep(attempt + 1)


def history_sheet(book):
    for sheet in book:
        if tuple(c.value for c in sheet[1][:12]) == HEADERS:
            return sheet
    raise LottoError("회차~총판매금액의 기존 12개 열 양식을 찾지 못했습니다.")


def read_draws(sheet):
    draws = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, max_col=12, values_only=True), 2):
        if all(value is None for value in values):
            continue
        draw = parse_draw(values)
        if draw.round != len(draws) + 1 or row_number != draw.round + 1:
            raise LottoError(f"{row_number}행: 회차는 1회부터 빈 행·중복 없이 오름차순이어야 합니다.")
        if draws and draw.drawn_at <= draws[-1].drawn_at:
            raise LottoError(f"{draw.round}회: 추첨일 순서가 잘못되었습니다.")
        draws.append(draw)
    if not draws:
        raise LottoError("엑셀에 분석할 당첨 이력이 없습니다.")
    return draws


def load_draws(path=DEFAULT_FILE):
    book = load_workbook(path, read_only=True)
    try:
        return read_draws(history_sheet(book))
    finally:
        book.close()


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save_atomic(book, path, expected_hash=None):
    """완성된 임시 파일을 검증하고 기존 파일을 백업한 뒤 한 번에 교체한다."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".tmp.xlsx", delete=False) as handle:
        temporary = Path(handle.name)
    backup = None
    try:
        book.save(temporary)
        check = load_workbook(temporary, read_only=True)
        try:
            for sheet in check:
                for _ in sheet.iter_rows(values_only=True):
                    pass
        finally:
            check.close()
        if expected_hash is not None and digest(path) != expected_hash:
            raise LottoError("작업 중 원본 엑셀이 변경되었습니다. 저장하지 않았으니 다시 실행하세요.")
        if path.exists():
            directory = path.parent / "backups"
            directory.mkdir(exist_ok=True)
            backup = directory / f"{path.stem}-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx"
            shutil.copy2(path, backup)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return backup


def update_history(path=DEFAULT_FILE, through=None, log=print):
    path = Path(path)
    original_hash = digest(path)
    book = load_workbook(path)
    try:
        sheet = history_sheet(book)
        existing = read_draws(sheet)
        latest = fetch_draw() if through is None else None
        target = latest.round if latest else integer(through, "마지막 회차", 1)
        incomplete = [d.round for d in existing if d.round <= target and None in d.cells()]
        needed = incomplete + list(range(existing[-1].round + 1, target + 1))
        if not needed:
            log(f"이미 {existing[-1].round}회까지 저장되어 있습니다. 변경 없음.")
            return {"added": 0, "repaired": 0, "latest": existing[-1].round, "backup": None}
        fetched = {}
        if latest is not None and latest.round in needed:
            fetched[latest.round] = latest
        for number in needed:
            draw = fetched.get(number) or fetch_draw(number)
            fetched[number] = draw
            if number <= len(existing):
                old = existing[number - 1].cells()
                for index, (previous, official) in enumerate(zip(old, draw.cells())):
                    if previous is not None and previous != official:
                        raise LottoError(f"{number}회 {HEADERS[index]}: 기존 값 {previous}와 공식 값 {official}가 다릅니다. 원본은 수정하지 않았습니다.")
            log(f"{number}회 공식 데이터 확인 ({len(fetched)}/{len(needed)})")
            time.sleep(0.15)
        for number in needed:
            for column, value in enumerate(fetched[number].cells(), 1):
                cell = sheet.cell(number + 1, column)
                if number > len(existing):
                    cell._style = copy(sheet.cell(len(existing) + 1, column)._style)
                if cell.value is None:
                    cell.value = value
        read_draws(sheet)
        backup = save_atomic(book, path, original_hash)
        added = sum(n > len(existing) for n in needed)
        log(f"완료: 새 회차 {added}개, 빈칸 보완 {len(incomplete)}개 회차. 백업: {backup}")
        return {"added": added, "repaired": len(incomplete), "latest": max(target, len(existing)), "backup": backup}
    finally:
        book.close()


def statistics(draws, window=52):
    window = integer(window, "최근 분석 회차 수", 1)
    if not draws:
        raise LottoError("분석할 당첨 이력이 없습니다.")
    recent = draws[-window:]
    total_counts = Counter(n for draw in draws for n in draw.numbers)
    recent_counts = Counter(n for draw in recent for n in draw.numbers)
    last_seen = {n: i for i, draw in enumerate(draws) for n in draw.numbers}
    return [{"number": n, "total": total_counts[n], "recent": recent_counts[n],
             "rate": total_counts[n] / len(draws),
             "gap": len(draws) - 1 - last_seen[n] if n in last_seen else len(draws)}
            for n in range(1, 46)]


def recommend(draws, games=2, mode="recent", window=52, rng=None):
    games = integer(games, "게임 수", 1)
    if games > 100 or mode not in MODES:
        raise LottoError("게임 수는 1~100이고, 추천 방식은 random/frequency/recent 중 하나여야 합니다.")
    stats = statistics(draws, window)
    rng = rng if rng is not None else random.SystemRandom()
    weights = [1 if mode == "random" else row["total" if mode == "frequency" else "recent"] + 1
               for row in stats]
    tickets = []
    seen = set()
    for _ in range(games * 100):
        available = list(range(1, 46))
        remaining = weights.copy()
        chosen = []
        for _ in range(6):
            n = rng.choices(available, weights=remaining, k=1)[0]
            index = available.index(n)
            available.pop(index)
            remaining.pop(index)
            chosen.append(n)
        ticket = tuple(sorted(chosen))
        if ticket not in seen:
            seen.add(ticket)
            tickets.append(ticket)
        if len(tickets) == games:
            return tickets
    raise LottoError("중복 없는 조합 생성에 실패했습니다. 다시 시도하세요.")


def rank(ticket, draw):
    matches = len(set(ticket) & set(draw.numbers))
    if matches == 6:
        return matches, 1
    if matches == 5:
        return matches, 2 if draw.bonus in ticket else 3
    return matches, {4: 4, 3: 5}.get(matches, 0)


def backtest(draws, rounds=104, games=2, window=52, seed=2026):
    rounds = integer(rounds, "검증 회차 수", 1)
    window = integer(window, "최근 분석 회차 수", 1)
    start = max(window, len(draws) - rounds)
    if start >= len(draws):
        raise LottoError("과거 검증에는 최근 분석 기간보다 많은 회차가 필요합니다.")
    results = []
    for mode in MODES:
        matches = Counter()
        ranks = Counter()
        for index in range(start, len(draws)):
            actual = draws[index]
            if actual.bonus is None:
                raise LottoError(f"{actual.round}회 보너스 번호가 없습니다. 엑셀 업데이트를 먼저 실행하세요.")
            # 검증 대상 회차와 이후 데이터는 추천에 절대 전달하지 않는다.
            tickets = recommend(draws[:index], games, mode, window,
                                random.Random(f"{seed}:{actual.round}:{mode}"))
            for ticket in tickets:
                count, prize_rank = rank(ticket, actual)
                matches[count] += 1
                ranks[prize_rank] += 1
        total = sum(matches.values())
        results.append({"mode": mode, "start": draws[start].round, "end": draws[-1].round,
                        "rounds": len(draws) - start, "tickets": total,
                        "average": sum(n * count for n, count in matches.items()) / total,
                        "matches": dict(matches), "ranks": dict(ranks), "seed": seed})
    return results


def style_report(book):
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243B53")
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) * 1.4 + 3
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(65, max(12, width))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=11)
                cell.alignment = Alignment(vertical="center", wrap_text=True)


def create_report(path=DEFAULT_FILE, games=2, mode="recent", window=52, seed=None,
                  test_rounds=104, output=None, log=print):
    path = Path(path)
    draws = load_draws(path)
    tickets = recommend(draws, games, mode, window, random.Random(seed) if seed is not None else None)
    stats = statistics(draws, window)
    log("과거 회차를 순서대로 검증합니다. 각 회차의 직전 데이터만 사용합니다.")
    results = backtest(draws, test_rounds, games, window)
    output = Path(output) if output else path.parent / "reports" / f"lotto-{draws[-1].round + 1}-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx"
    if output.resolve() == path.resolve():
        raise LottoError("분석 결과는 원본 lotto.xlsx와 다른 파일에 저장해야 합니다.")
    book = Workbook()
    summary = book.active
    summary.title = "안내"
    for row in [("항목", "내용"), ("안내", NOTICE), ("원본", str(path.resolve())),
                ("사용 데이터", f"1~{draws[-1].round}회 ({draws[-1].drawn_at.isoformat()} 추첨까지)"),
                ("추천 대상", f"{draws[-1].round + 1}회 — 원본의 마지막 회차 다음 회차"),
                ("추천 방식", MODES[mode]), ("최근 분석 기간", min(window, len(draws))),
                ("추천 시드", seed if seed is not None else "매번 새로운 난수"),
                ("추천 규칙", "빈도 참고 방식은 출현 횟수+1을 가중치로 6개를 중복 없이 추출합니다. 무작위 방식은 모든 번호의 가중치가 같습니다."),
                ("통계 범위", "보너스는 출현 통계에서 제외하며, 2등 판정에만 사용합니다."),
                ("과거 검증", "각 회차 직전의 이력만 사용합니다. 같은 게임 수로 세 방식을 비교하며 검증 시드는 2026으로 고정합니다."),
                ("검증 한계", "한 번의 모의실험입니다. 차이는 우연일 수 있으며 통계적 유의성이나 미래 예측력을 입증하지 않습니다."),
                ("이론상 평균 일치", "공정한 추첨에서 게임당 0.8개 (6×6÷45)"),
                ("만든 시각", datetime.now(KST).isoformat(timespec="seconds")), ("공식 API", API)]:
        summary.append(row)
    recommendations = book.create_sheet("추천번호")
    recommendations.append(["대상회차", "게임", "방식", *HEADERS[2:8]])
    for index, ticket in enumerate(tickets, 1):
        recommendations.append([draws[-1].round + 1, index, MODES[mode], *ticket])
    counts = book.create_sheet("번호통계")
    counts.append(["번호", "전체 출현 횟수", f"최근 {min(window, len(draws))}회 출현 횟수", "회차당 출현 비율", "연속 미출현 회차"])
    for row in stats:
        counts.append([row["number"], row["total"], row["recent"], row["rate"], row["gap"]])
        counts.cell(counts.max_row, 4).number_format = "0.00%"
    chart = BarChart()
    chart.title = "번호별 전체 출현 횟수 (보너스 제외)"
    chart.add_data(Reference(counts, min_col=2, min_row=1, max_row=46), titles_from_data=True)
    chart.set_categories(Reference(counts, min_col=1, min_row=2, max_row=46))
    chart.width, chart.height = 30, 12
    counts.add_chart(chart, "G2")
    verification = book.create_sheet("과거검증")
    verification.append(["방식", "시작회차", "종료회차", "검증 회차 수", "게임 수", "평균 일치 개수",
                         "1등", "2등", "3등", "4등", "5등", *[f"{n}개 일치" for n in range(7)]])
    for result in results:
        verification.append([MODES[result["mode"]], result["start"], result["end"], result["rounds"],
                             result["tickets"], result["average"],
                             *[result["ranks"].get(n, 0) for n in range(1, 6)],
                             *[result["matches"].get(n, 0) for n in range(7)]])
        verification.cell(verification.max_row, 6).number_format = "0.0000"
    style_report(book)
    summary.column_dimensions["B"].width = 100
    for row in range(2, summary.max_row + 1):
        summary.row_dimensions[row].height = 44
    try:
        save_atomic(book, output)
    finally:
        book.close()
    log(f"{draws[-1].round + 1}회 참고용 추천 ({MODES[mode]}):")
    for index, ticket in enumerate(tickets, 1):
        log(f"  {index}: {'  '.join(f'{n:02}' for n in ticket)}")
    for result in results:
        log(f"과거 검증 · {MODES[result['mode']]}: 게임당 평균 {result['average']:.4f}개 일치 / {result['tickets']}게임")
    log(f"결과 저장: {output}")
    log(NOTICE)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", type=Path, default=DEFAULT_FILE, help="원본 엑셀 경로")
    commands = parser.add_subparsers(dest="command", required=True)
    update = commands.add_parser("update", help="공식 API로 새 회차와 빈칸 갱신")
    update.add_argument("--to", type=int, help="생략하면 API의 최신 확정 회차까지")
    weekly = commands.add_parser("weekly", help="심층 분석 후 이번 회차의 고정 2게임 저장")
    weekly.add_argument("--offline", action="store_true", help="최신 회차 확인 없이 저장된 이력만 분석")
    report = commands.add_parser("report", help="분석·추천·과거 검증을 새 엑셀에 저장")
    report.add_argument("--games", type=int, default=2)
    report.add_argument("--window", type=int, default=52)
    report.add_argument("--mode", choices=MODES, default="recent")
    report.add_argument("--seed", type=int)
    report.add_argument("--test-rounds", type=int, default=104)
    report.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        if args.command == "update":
            update_history(args.file, args.to)
        elif args.command == "weekly":
            from lotto_weekly import weekly
            weekly(args.file, offline=args.offline)
        else:
            create_report(args.file, args.games, args.mode, args.window, args.seed, args.test_rounds, args.output)
    except (LottoError, OSError) as exc:
        parser.exit(1, f"오류: {exc}\n엑셀이 열려 있으면 닫고 다시 실행하세요.\n")


if __name__ == "__main__":
    main()
